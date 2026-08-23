"""Local document storage, durable indexing jobs, and lexical retrieval."""

from __future__ import annotations

import csv
import io
import re
import sqlite3
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from uuid import uuid4

from app.schemas.domain import KnowledgeHit

SUPPORTED_TYPES = {"TXT": ".txt", "MD": ".md", "PDF": ".pdf", "DOCX": ".docx", "CSV": ".csv"}
SUPPORTED_MIME_TYPES = {
    "TXT": {"text/plain"},
    "MD": {"text/markdown", "text/plain"},
    "CSV": {"text/csv", "text/plain", "application/vnd.ms-excel"},
    "PDF": {"application/pdf"},
    "DOCX": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
}
MAX_CATEGORY_LENGTH = 80
MAX_SUMMARY_LENGTH = 240
MAX_EXCERPT_LENGTH = 420
JOB_LOCK_TIMEOUT = timedelta(minutes=15)
TOKEN_PATTERN = re.compile(r"[a-z0-9_]+|[\u4e00-\u9fff]", re.IGNORECASE)


class KnowledgeBusyError(RuntimeError):
    """Raised when a document cannot be deleted while it is being indexed."""


class KnowledgeStore:
    def __init__(
        self,
        database_path: str | Path,
        root: str | Path,
        *,
        max_upload_bytes: int = 20 * 1024 * 1024,
        top_k: int = 5,
        chunk_size: int = 800,
        chunk_overlap: int = 120,
        workers: int = 2,
    ) -> None:
        self.database_path = Path(database_path).expanduser()
        self.root = Path(root).expanduser()
        self.max_upload_bytes = max_upload_bytes
        self.top_k = top_k
        self.chunk_size = max(100, chunk_size)
        self.chunk_overlap = min(max(0, chunk_overlap), self.chunk_size // 2)
        self.root.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._submitted: set[str] = set()
        self._executor = ThreadPoolExecutor(max_workers=max(1, workers), thread_name_prefix="knowledge")
        self._initialize()
        self._recover_jobs()
        self._submit_queued()

    @contextmanager
    def _connection(self):
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        connection = sqlite3.connect(self.database_path, timeout=10, check_same_thread=False)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        connection.execute("PRAGMA busy_timeout = 10000")
        try:
            yield connection
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            connection.close()

    def _initialize(self) -> None:
        with self._connection() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_documents (
                    id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    file_type TEXT NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    category TEXT NOT NULL,
                    status TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    summary TEXT NOT NULL DEFAULT '',
                    failure_message TEXT,
                    original_path TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS knowledge_chunks (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL REFERENCES knowledge_documents(id) ON DELETE CASCADE,
                    ordinal INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    UNIQUE(document_id, ordinal)
                );
                CREATE TABLE IF NOT EXISTS knowledge_jobs (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL UNIQUE REFERENCES knowledge_documents(id) ON DELETE CASCADE,
                    status TEXT NOT NULL,
                    attempts INTEGER NOT NULL DEFAULT 0,
                    locked_at TEXT,
                    error TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_documents_status ON knowledge_documents(status);
                CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_document ON knowledge_chunks(document_id, ordinal);
                """
            )
            try:
                connection.execute(
                    """CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_chunks_fts USING fts5(
                        chunk_id UNINDEXED, document_id UNINDEXED, content
                    )"""
                )
            except sqlite3.OperationalError:
                # Retrieval has a deterministic Python fallback when FTS5 is unavailable.
                pass

    def _recover_jobs(self) -> None:
        with self._connection() as connection:
            now = _now()
            running = connection.execute(
                "SELECT id, document_id, attempts, locked_at FROM knowledge_jobs WHERE status = 'running'"
            ).fetchall()
            for row in running:
                if row["attempts"] >= 3:
                    connection.execute(
                        "UPDATE knowledge_jobs SET status = 'failed', locked_at = NULL, error = COALESCE(error, '任务重试次数已用尽。'), updated_at = ? WHERE id = ?",
                        (now, row["id"]),
                    )
                    connection.execute(
                        "UPDATE knowledge_documents SET status = 'failed', failure_message = COALESCE(failure_message, '任务重试次数已用尽。'), updated_at = ? WHERE id = ?",
                        (now, row["document_id"]),
                    )
                elif _lock_expired(row["locked_at"]):
                    connection.execute(
                        "UPDATE knowledge_jobs SET status = 'queued', locked_at = NULL, updated_at = ? WHERE id = ?",
                        (now, row["id"]),
                    )
            connection.execute(
                "UPDATE knowledge_jobs SET status = 'failed', locked_at = NULL, error = COALESCE(error, '任务重试次数已用尽。'), updated_at = ? WHERE status = 'queued' AND attempts >= 3",
                (now,),
            )
            connection.execute(
                "UPDATE knowledge_documents SET status = 'failed', failure_message = COALESCE(failure_message, '任务重试次数已用尽。'), updated_at = ? WHERE id IN (SELECT document_id FROM knowledge_jobs WHERE status = 'failed' AND attempts >= 3)",
                (now,),
            )

    def _submit_queued(self) -> None:
        with self._connection() as connection:
            ids = [row["id"] for row in connection.execute("SELECT id FROM knowledge_jobs WHERE status = 'queued'")]
        for job_id in ids:
            self._submit(job_id)

    def _submit(self, job_id: str) -> None:
        with self._lock:
            if job_id in self._submitted:
                return
            self._submitted.add(job_id)
        self._executor.submit(self._run_job, job_id)

    def create_upload(self, filename: str, content: bytes, category: str, content_type: str | None = None) -> dict[str, Any]:
        if "\x00" in (filename or ""):
            raise ValueError("文件名无效。")
        safe_filename = Path(filename or "").name.strip() or "未命名文档"
        suffix = Path(safe_filename).suffix.lower()
        file_type = next((kind for kind, extension in SUPPORTED_TYPES.items() if extension == suffix), None)
        if file_type is None:
            raise ValueError("仅支持 TXT、MD、PDF、DOCX 和 CSV 文件。")
        if not content:
            raise ValueError("不能上传空文件。")
        if len(content) > self.max_upload_bytes:
            raise ValueError(f"文件大小不能超过 {self.max_upload_bytes // 1024 // 1024} MiB。")
        _validate_signature(file_type, content, content_type)
        normalized_category = (category or "未分类").strip()[:MAX_CATEGORY_LENGTH] or "未分类"
        document_id = str(uuid4())
        now = _now()
        original_dir = self.root / "originals"
        original_dir.mkdir(parents=True, exist_ok=True)
        path = original_dir / f"{document_id}{suffix}"
        path.write_bytes(content)
        try:
            with self._connection() as connection:
                connection.execute(
                    """INSERT INTO knowledge_documents
                    (id, filename, file_type, size_bytes, category, status, created_at, updated_at, original_path)
                    VALUES (?, ?, ?, ?, ?, 'uploading', ?, ?, ?)""",
                    (document_id, safe_filename[:255], file_type, len(content), normalized_category, now, now, str(path)),
                )
                connection.execute(
                    """INSERT INTO knowledge_jobs
                    (id, document_id, status, created_at, updated_at) VALUES (?, ?, 'queued', ?, ?)""",
                    (str(uuid4()), document_id, now, now),
                )
        except Exception:
            path.unlink(missing_ok=True)
            raise
        response = {
            "id": document_id,
            "filename": safe_filename[:255],
            "file_type": file_type,
            "size_bytes": len(content),
            "category": normalized_category,
            "status": "uploading",
            "created_at": now,
            "updated_at": now,
            "chunk_count": 0,
            "summary": "",
            "failure_message": None,
        }
        self._submit_queued()
        return response

    def list(self) -> list[dict[str, Any]]:
        with self._connection() as connection:
            rows = connection.execute(
                "SELECT id, filename, file_type, size_bytes, category, status, created_at, updated_at, "
                "chunk_count, summary, failure_message FROM knowledge_documents ORDER BY created_at DESC"
            ).fetchall()
        return [dict(row) for row in rows]

    def get(self, document_id: str) -> dict[str, Any] | None:
        with self._connection() as connection:
            row = connection.execute(
                "SELECT id, filename, file_type, size_bytes, category, status, created_at, updated_at, "
                "chunk_count, summary, failure_message FROM knowledge_documents WHERE id = ?",
                (document_id,),
            ).fetchone()
        return dict(row) if row else None

    def delete(self, document_id: str) -> None:
        with self._connection() as connection:
            row = connection.execute(
                "SELECT status, original_path FROM knowledge_documents WHERE id = ?", (document_id,)
            ).fetchone()
            if row is None:
                raise KeyError(document_id)
            if row["status"] in {"uploading", "parsing"}:
                raise KnowledgeBusyError("文档正在解析或索引，请稍后再删除。")
            try:
                connection.execute("DELETE FROM knowledge_chunks_fts WHERE document_id = ?", (document_id,))
            except sqlite3.OperationalError:
                pass
            connection.execute("DELETE FROM knowledge_chunks WHERE document_id = ?", (document_id,))
            connection.execute("DELETE FROM knowledge_documents WHERE id = ?", (document_id,))
        try:
            Path(row["original_path"]).unlink(missing_ok=True)
        except OSError:
            pass

    def retrieve(self, question: str, top_k: int | None = None) -> list[KnowledgeHit]:
        limit = max(1, top_k or self.top_k)
        if not question.strip():
            return []
        with self._connection() as connection:
            rows = self._fts_rows(connection, question, limit)
            if rows is None:
                rows = connection.execute(
                    "SELECT c.id, c.document_id, c.content, d.filename, d.category "
                    "FROM knowledge_chunks c JOIN knowledge_documents d ON d.id = c.document_id "
                    "WHERE d.status = 'indexed' ORDER BY d.created_at DESC"
                ).fetchall()
        if not rows:
            return []
        query_tokens = _tokens(question)
        scored: list[tuple[float, sqlite3.Row]] = []
        for row in rows:
            if "rank" in row.keys():
                rank = float(row["rank"])
                # SQLite's BM25 returns lower (usually more negative) values for
                # better matches; map that ordering onto the public 0..1 score.
                strength = max(0.0, -rank)
                scored.append((strength / (1.0 + strength), row))
                continue
            tokens = _tokens(f"{row['filename']} {row['category']} {row['content']}")
            overlap = len(set(query_tokens).intersection(tokens))
            if overlap:
                score = float(overlap) / max(1, len(set(query_tokens)))
                scored.append((score, row))
        scored.sort(key=lambda item: (-item[0], item[1]["filename"], item[1]["id"]))
        hits: list[KnowledgeHit] = []
        seen: set[str] = set()
        for score, row in scored:
            if row["document_id"] in seen:
                continue
            seen.add(row["document_id"])
            hits.append(
                KnowledgeHit(
                    document_id=row["document_id"],
                    title=row["filename"],
                    category=row["category"],
                    excerpt=_excerpt(row["content"]),
                    relevance=round(min(1.0, score), 4),
                )
            )
            if len(hits) >= limit:
                break
        return hits

    @staticmethod
    def _fts_rows(connection: sqlite3.Connection, question: str, limit: int) -> list[sqlite3.Row] | None:
        """Return BM25-ranked rows, or None when this SQLite build has no FTS5 table."""
        terms = _tokens(question)
        if not terms:
            return []
        match_query = " OR ".join(f'"{term.replace(chr(34), chr(34) * 2)}"' for term in dict.fromkeys(terms))
        try:
            return connection.execute(
                "SELECT f.chunk_id AS id, f.document_id, f.content, d.filename, d.category, bm25(f) AS rank "
                "FROM knowledge_chunks_fts f JOIN knowledge_documents d ON d.id = f.document_id "
                "WHERE knowledge_chunks_fts MATCH ? AND d.status = 'indexed' "
                "ORDER BY rank LIMIT ?",
                (match_query, max(limit * 8, limit)),
            ).fetchall()
        except sqlite3.OperationalError:
            return None

    def close(self) -> None:
        self._executor.shutdown(wait=False, cancel_futures=True)

    def _run_job(self, job_id: str) -> None:
        try:
            claimed = self._claim_job(job_id)
            if not claimed:
                return
            document = self.get(claimed)
            if not document:
                return
            self._set_document_status(claimed, "parsing")
            text = _extract_text(document["file_type"], Path(self._original_path(claimed)))
            if not text.strip():
                raise ValueError("文档没有可索引的文本内容。")
            if len(text) > 2_000_000:
                raise ValueError("解析后的文本超过 2,000,000 个字符。")
            chunks = _chunk_text(text, self.chunk_size, self.chunk_overlap)
            summary = _excerpt(text, MAX_SUMMARY_LENGTH)
            now = _now()
            with self._connection() as connection:
                connection.execute("DELETE FROM knowledge_chunks WHERE document_id = ?", (claimed,))
                try:
                    connection.execute("DELETE FROM knowledge_chunks_fts WHERE document_id = ?", (claimed,))
                except sqlite3.OperationalError:
                    pass
                for ordinal, chunk in enumerate(chunks):
                    chunk_id = str(uuid4())
                    connection.execute(
                        "INSERT INTO knowledge_chunks(id, document_id, ordinal, content) VALUES (?, ?, ?, ?)",
                        (chunk_id, claimed, ordinal, chunk),
                    )
                    try:
                        connection.execute(
                            "INSERT INTO knowledge_chunks_fts(chunk_id, document_id, content) VALUES (?, ?, ?)",
                            (chunk_id, claimed, chunk),
                        )
                    except sqlite3.OperationalError:
                        pass
                connection.execute(
                    "UPDATE knowledge_documents SET status = 'indexed', updated_at = ?, chunk_count = ?, summary = ?, failure_message = NULL WHERE id = ?",
                    (now, len(chunks), summary, claimed),
                )
                connection.execute(
                    "UPDATE knowledge_jobs SET status = 'done', locked_at = NULL, error = NULL, updated_at = ? WHERE document_id = ?",
                    (now, claimed),
                )
        except Exception as error:
            self._fail_job(job_id, str(error))
        finally:
            with self._lock:
                self._submitted.discard(job_id)

    def _claim_job(self, job_id: str) -> str | None:
        with self._connection() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT document_id, status, attempts FROM knowledge_jobs WHERE id = ?", (job_id,)
            ).fetchone()
            if not row or row["status"] != "queued":
                return None
            if row["attempts"] >= 3:
                now = _now()
                connection.execute(
                    "UPDATE knowledge_jobs SET status = 'failed', locked_at = NULL, error = COALESCE(error, '任务重试次数已用尽。'), updated_at = ? WHERE id = ?",
                    (now, job_id),
                )
                connection.execute(
                    "UPDATE knowledge_documents SET status = 'failed', failure_message = COALESCE(failure_message, '任务重试次数已用尽。'), updated_at = ? WHERE id = ?",
                    (now, row["document_id"]),
                )
                return None
            now = _now()
            connection.execute(
                "UPDATE knowledge_jobs SET status = 'running', attempts = attempts + 1, locked_at = ?, updated_at = ? WHERE id = ?",
                (now, now, job_id),
            )
            return row["document_id"]

    def _fail_job(self, job_id: str, error: str) -> None:
        message = error[:500] or "文档解析失败。"
        with self._connection() as connection:
            row = connection.execute("SELECT document_id, attempts FROM knowledge_jobs WHERE id = ?", (job_id,)).fetchone()
            if not row:
                return
            final = row["attempts"] >= 3
            status = "failed" if final else "queued"
            now = _now()
            connection.execute(
                "UPDATE knowledge_jobs SET status = ?, locked_at = NULL, error = ?, updated_at = ? WHERE id = ?",
                (status, message, now, job_id),
            )
            connection.execute(
                "UPDATE knowledge_documents SET status = ?, updated_at = ?, failure_message = ? WHERE id = ?",
                ("failed" if final else "uploading", now, message, row["document_id"]),
            )
        if not final:
            with self._lock:
                self._submitted.discard(job_id)
            self._submit(job_id)

    def _set_document_status(self, document_id: str, status: str) -> None:
        with self._connection() as connection:
            connection.execute("UPDATE knowledge_documents SET status = ?, updated_at = ? WHERE id = ?", (status, _now(), document_id))

    def _original_path(self, document_id: str) -> str:
        with self._connection() as connection:
            row = connection.execute("SELECT original_path FROM knowledge_documents WHERE id = ?", (document_id,)).fetchone()
        if not row:
            raise FileNotFoundError(document_id)
        return row["original_path"]


_stores: dict[str, KnowledgeStore] = {}
_stores_lock = threading.Lock()


def get_knowledge_store(settings: Any) -> KnowledgeStore:
    key = str(Path(settings.knowledge_database_path).expanduser())
    with _stores_lock:
        store = _stores.get(key)
        if store is None:
            store = KnowledgeStore(
                settings.knowledge_database_path,
                settings.knowledge_root,
                max_upload_bytes=settings.knowledge_max_upload_bytes,
                top_k=settings.knowledge_top_k,
                chunk_size=settings.knowledge_chunk_size,
                chunk_overlap=settings.knowledge_chunk_overlap,
            )
            _stores[key] = store
        return store


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _lock_expired(value: str | None) -> bool:
    if not value:
        return True
    try:
        locked_at = datetime.fromisoformat(value)
    except (TypeError, ValueError):
        return True
    if locked_at.tzinfo is None:
        locked_at = locked_at.replace(tzinfo=UTC)
    return datetime.now(UTC) - locked_at >= JOB_LOCK_TIMEOUT


def _tokens(value: str) -> list[str]:
    return TOKEN_PATTERN.findall(value.lower())


def _excerpt(value: str, limit: int = MAX_EXCERPT_LENGTH) -> str:
    text = re.sub(r"\s+", " ", value).strip()
    return text[:limit] + ("…" if len(text) > limit else "")


def _chunk_text(text: str, size: int, overlap: int) -> list[str]:
    normalized = text.replace("\r\n", "\n").strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    step = max(1, size - overlap)
    while start < len(normalized):
        chunk = normalized[start : start + size].strip()
        if chunk:
            chunks.append(chunk)
        start += step
    return chunks


def _validate_signature(file_type: str, content: bytes, content_type: str | None) -> None:
    if file_type == "PDF" and not content.startswith(b"%PDF"):
        raise ValueError("PDF 文件签名无效。")
    if file_type == "DOCX":
        if not content.startswith(b"PK"):
            raise ValueError("DOCX 文件签名无效。")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
        except (zipfile.BadZipFile, OSError) as error:
            raise ValueError("DOCX 文件签名无效。") from error
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ValueError("DOCX 文件签名无效。")
    if content_type:
        normalized_type = content_type.split(";", 1)[0].strip().lower()
        if normalized_type != "application/octet-stream" and normalized_type not in SUPPORTED_MIME_TYPES[file_type]:
            raise ValueError("文件类型不受支持。")


def _extract_text(file_type: str, path: Path) -> str:
    content = path.read_bytes()
    if file_type in {"TXT", "MD"}:
        return content.decode("utf-8-sig", errors="replace")
    if file_type == "CSV":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig", errors="replace")))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    if file_type == "PDF":
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise RuntimeError("PDF 解析依赖未安装，请安装 pypdf。") from error
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if file_type == "DOCX":
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("DOCX 解析依赖未安装，请安装 python-docx。") from error
        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    raise ValueError("文件类型不受支持。")
